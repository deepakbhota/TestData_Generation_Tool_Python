"""Synthetic sample files used by the tests (and handy for trying the UI)."""
from __future__ import annotations

import io

PIPE_V5 = (
    "HDR|20240105|PAYMENTS\n"
    "MANUFACTURER|PAY_DATE|AMOUNT|DC_IND|PAYMENT_REF\n"
    "AUD|20240104|1500.25|CR|PAY000101\n"
    "AUD|20240104|200.00|DR|PAY000102\n"
    "BMW|20240104|9800.10|CR|PAY000103\n"
    "BMW|20240104|450.55|CR|PAY000104\n"
    "TRL|4|11550.90\n"
)

FIXED_V5 = (
    "H 20240105 PAYFILE\n"
    "AUD 20240104 000000150025 C PAY000101\n"
    "AUD 20240104 000000020000 D PAY000102\n"
    "BMW 20240104 000000980010 C PAY000103\n"
    "BMW 20240104 000000045055 C PAY000104\n"
    "T 0000004\n"
)

CSV_DD = (
    '"JLR","05/01/2024","0001","1250.00"\n'
    '"JLR","04/01/2024","0002","310.40"\n'
    '"JLR","03/01/2024","0003","99.99"\n'
)

CSV_DC = (
    "Manufacturer,Payment Date,Credit Amount,Debit Amount,Invoice\r\n"
    "JLR,04/01/2024,5000.00,0.00,INV-100001\r\n"
    "JLR,04/01/2024,1200.50,0.00,INV-100002\r\n"
    "JLR,03/01/2024,300.00,0.00,INV-100003\r\n"
)


def make_xlsx() -> bytes:
    from datetime import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Payments"
    ws.append(["Payment file", datetime(2024, 1, 5)])
    ws.append(["Manufacturer", "Payment Date", "Amount", "Type", "Reference"])
    for c in ws[2]:
        c.font = Font(bold=True)
    rows = [
        ("AUD", datetime(2024, 1, 4), 1500.25, "CR", "R0001"),
        ("AUD", datetime(2024, 1, 4), 200.0, "DR", "R0002"),
        ("BMW", datetime(2024, 1, 4), 9800.1, "CR", "R0003"),
    ]
    for r in rows:
        ws.append(r)
    for r in range(3, 6):
        ws.cell(row=r, column=2).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=3).number_format = "#,##0.00"
    ws.append(["Total rows", 3, 11500.35])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
