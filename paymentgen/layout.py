"""Infer a payment-file layout (schema) from an uploaded sample file.

Supported samples: delimited text (csv/txt/dat/...), fixed-width text and Excel
(.xlsx/.xlsm).  A sample is split into optional *header* records, *detail*
records and optional *trailer* records.  Every detail column gets a ``role``
(manufacturer, payment date, amount, ...) which the generator uses to fill in
values; roles are auto-detected and can be overridden by the user.
"""
from __future__ import annotations

import copy
import io
import os
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

ROLES = [
    "manufacturer",
    "payment_date",
    "file_date",
    "amount",
    "credit_amount",
    "debit_amount",
    "dc_indicator",
    "sequence",
    "reference",
    "other",
]

# (strptime format, regex that finds such a date inside free text)
DATE_FORMATS: list[tuple[str, str]] = [
    ("%Y%m%d", r"(?:19|20)\d{2}(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])"),
    ("%Y-%m-%d", r"\d{4}-\d{2}-\d{2}"),
    ("%Y/%m/%d", r"\d{4}/\d{2}/\d{2}"),
    ("%d/%m/%Y", r"\d{2}/\d{2}/\d{4}"),
    ("%d-%m-%Y", r"\d{2}-\d{2}-\d{4}"),
    ("%d.%m.%Y", r"\d{2}\.\d{2}\.\d{4}"),
    ("%d-%b-%Y", r"\d{2}-[A-Za-z]{3}-\d{4}"),
    ("%d%b%Y", r"\d{2}[A-Za-z]{3}\d{4}"),
    ("%d/%m/%y", r"\d{2}/\d{2}/\d{2}"),
    ("%d%m%Y", r"\d{8}"),
]
_DEBIT_TOKENS = {"D", "DR", "DB", "DEBIT"}
_CREDIT_TOKENS = {"C", "CR", "CRD", "CREDIT"}
_DELIMITERS = [",", "|", ";", "\t"]
CENT = Decimal("0.01")


class SampleError(ValueError):
    """The sample file cannot be understood."""


# --------------------------------------------------------------------------- model
@dataclass
class Column:
    name: str
    role: str = "auto"
    start: int = 0  # fixed-width only
    width: int = 0  # fixed-width only
    kind: str = "text"  # text | int | decimal | date
    native: bool = False  # Excel: cell holds a real number/date, not text
    decimals: int = 0
    implied_point: bool = False  # amount stored as integer minor units, e.g. 0000012345
    zero_pad: int = 0
    thousands: bool = False
    date_fmt: str = ""
    date_upper: bool = False
    align: str = "left"
    quoted: bool = False  # delimited: sample wrapped this column in quotes
    blank_when_unused: bool = False
    pool: list[str] = field(default_factory=list)
    credit_token: str = "CR"
    debit_token: str = "DR"
    prefix: str = ""  # reference columns
    digits: int = 0
    number_format: str = "General"


@dataclass
class Binding:
    """A header/trailer value that must be refreshed (record count, total ...)."""

    section: str  # "header" | "trailer"
    line: int
    stat: str  # count | credit_sum | debit_sum | total_sum | net
    index: int = -1  # field / cell index (delimited, excel)
    start: int = -1  # character span (fixed)
    end: int = -1
    decimals: int = 0
    has_dot: bool = False
    pad: int = 0


@dataclass
class Layout:
    source_name: str
    ext: str
    fmt: str  # delimited | fixed | excel
    delimiter: str = ","
    encoding: str = "utf-8"
    newline: str = "\n"
    final_newline: bool = True
    sheet_name: str = "Sheet1"
    header: list[Any] = field(default_factory=list)  # str lines or Excel rows
    trailer: list[Any] = field(default_factory=list)
    details: list[Any] = field(default_factory=list)
    has_column_header: bool = False
    columns: list[Column] = field(default_factory=list)
    line_length: int = 0
    debit_style: str = "unsigned"  # indicator | columns | signed | unsigned
    bindings: list[Binding] = field(default_factory=list)
    cell_styles: dict = field(default_factory=dict)
    col_widths: dict = field(default_factory=dict)

    # -- raw record access ---------------------------------------------------
    def fields(self, record: Any) -> list[Any]:
        """Split one raw record into field values (parallel to ``columns``)."""
        if self.fmt == "excel":
            return list(record)
        if self.fmt == "delimited":
            return [v for v, _ in split_delimited(record, self.delimiter)]
        line = record.ljust(self.line_length)
        return [line[c.start : c.start + c.width].strip() for c in self.columns]

    def column_values(self, index: int) -> list[Any]:
        out = []
        for rec in self.details:
            vals = self.fields(rec)
            out.append(vals[index] if index < len(vals) else None)
        return out

    @property
    def column_names(self) -> list[str]:
        return [c.name for c in self.columns]

    def role_column(self, role: str) -> Column | None:
        return next((c for c in self.columns if c.role == role), None)

    def warnings(self, debits_needed: bool = False) -> list[str]:
        out = []
        roles = {c.role for c in self.columns}
        if "manufacturer" not in roles:
            out.append("No column is mapped to **manufacturer** - generated rows will not carry a manufacturer code.")
        if not roles & {"amount", "credit_amount", "debit_amount"}:
            out.append("No amount column found - map at least one column to **amount**.")
        if "payment_date" not in roles:
            out.append("No **payment_date** column - payment dates of the covered days will not appear in the rows.")
        if debits_needed and self.debit_style == "unsigned":
            out.append(
                "The sample does not show how debits are represented (no DR/CR indicator, no separate debit "
                "column and no negative amounts). Debit rows will be written as negative amounts."
            )
        return out


# --------------------------------------------------------------------------- helpers
def split_delimited(line: str, delim: str) -> list[tuple[str, bool]]:
    """Split a line into (value, was_quoted) pairs, honouring ``"`` quoting."""
    out: list[tuple[str, bool]] = []
    i, n = 0, len(line)
    while True:
        if i < n and line[i] == '"':
            j, buf = i + 1, []
            while j < n:
                if line[j] == '"':
                    if j + 1 < n and line[j + 1] == '"':
                        buf.append('"')
                        j += 2
                        continue
                    break
                buf.append(line[j])
                j += 1
            k = line.find(delim, j)
            out.append(("".join(buf), True))
            i = n if k == -1 else k
        else:
            k = line.find(delim, i)
            end = n if k == -1 else k
            out.append((line[i:end], False))
            i = end
        if i >= n:
            break
        i += 1
        if i == n:
            out.append(("", False))
            break
    return out


def join_delimited(values: list[str], quoted: list[bool], delim: str) -> str:
    parts = []
    for v, q in zip(values, quoted):
        if q or any(ch in v for ch in (delim, '"', "\n", "\r")):
            v = '"' + v.replace('"', '""') + '"'
        parts.append(v)
    return delim.join(parts)


def parse_date_str(s: str, fmt: str) -> date | None:
    try:
        d = datetime.strptime(s, fmt)
    except ValueError:
        return None
    if d.strftime(fmt).lower() != s.lower() or not 1990 <= d.year <= 2100:
        return None
    return d.date()


def detect_date_format(values: list[str]) -> str | None:
    if not values:
        return None
    for fmt, _ in DATE_FORMATS:
        if all(parse_date_str(v, fmt) for v in values):
            return fmt
    return None


def refresh_dates(text: str, new_date: date) -> str:
    """Replace every date found in free text (header/trailer) by ``new_date``."""
    for fmt, rx in DATE_FORMATS:
        pattern = re.compile(rf"(?<![\d A-Za-z]){rx}(?![\d])" if fmt == "%d%m%Y" else rf"(?<!\d){rx}(?!\d)")

        def sub(m: re.Match, fmt=fmt) -> str:
            if not parse_date_str(m.group(0), fmt):
                return m.group(0)
            out = new_date.strftime(fmt)
            return out.upper() if m.group(0).isupper() else out

        text = pattern.sub(sub, text)
    return text


def _is_num(v: Any) -> bool:
    return isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)


def _numlike(v: Any) -> bool:
    if _is_num(v) or isinstance(v, (datetime, date)):
        return True
    s = str(v).strip().replace(",", "")
    return bool(re.fullmatch(r"-?\d+(\.\d+)?", s)) or detect_date_format([s]) is not None


def _blank(v: Any) -> bool:
    return v is None or str(v).strip() == ""


def to_decimal(raw: Any, col: Column | None = None) -> Decimal | None:
    if _blank(raw):
        return None
    try:
        d = Decimal(str(raw).strip().replace(",", ""))
    except InvalidOperation:
        return None
    if col is not None and col.implied_point and not (isinstance(raw, str) and "." in raw):
        d = d / (Decimal(10) ** col.decimals)
    return d


def _decode(data: bytes) -> tuple[str, str]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig"), "utf-8-sig"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return data.decode("latin-1"), "latin-1"


def _dedupe(names: list[str]) -> list[str]:
    seen: Counter = Counter()
    out = []
    for n in names:
        seen[n] += 1
        out.append(n if seen[n] == 1 else f"{n}_{seen[n]}")
    return out


# --------------------------------------------------------------------------- parsing
def parse_sample(filename: str, data: bytes) -> Layout:
    """Read a sample file and return its auto-detected :class:`Layout`."""
    ext = os.path.splitext(filename)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        layout = _parse_excel(filename, ext, data)
    elif ext in (".xls", ".xlsb"):
        raise SampleError("Legacy Excel formats are not supported - please save the sample as .xlsx.")
    else:
        layout = _parse_text(filename, ext, data)
    analyze(layout)
    return layout


def _peel(keys: list[Any]) -> tuple[int, int]:
    """Number of leading header / trailing trailer records."""
    n = len(keys)
    if n < 3:
        return 0, 0
    modal = Counter(keys[1:-1]).most_common(1)[0][0]
    idx = [i for i, k in enumerate(keys) if k == modal]
    return idx[0], n - 1 - idx[-1]


def _looks_like_header(first: list[Any], rest: list[list[Any]]) -> bool:
    cells = [c for c in first if not _blank(c)]
    if not cells or any(_numlike(c) for c in cells):
        return False
    if len({str(c).strip() for c in cells}) != len(cells):
        return False
    for ci in range(len(first)):
        col = [r[ci] for r in rest if ci < len(r) and not _blank(r[ci])]
        if col and all(_numlike(v) for v in col):
            return True
    has_digit_rest = any(ch.isdigit() for r in rest for c in r for ch in str(c))
    return has_digit_rest and not any(ch.isdigit() for c in cells for ch in str(c))


def _detect_format(lines: list[str]) -> tuple[str, str]:
    sample = lines[:200]
    best = None
    for d in _DELIMITERS:
        counts = [len(split_delimited(l, d)) for l in sample]
        modal, freq = Counter(counts).most_common(1)[0]
        frac = freq / len(counts)
        if modal >= 2 and frac >= 0.6 and (best is None or (frac, modal) > best[0]):
            best = ((frac, modal), d)
    return ("delimited", best[1]) if best else ("fixed", "")


def _parse_text(filename: str, ext: str, data: bytes) -> Layout:
    text, enc = _decode(data)
    newline = "\r\n" if "\r\n" in text else ("\r" if "\r" in text and "\n" not in text else "\n")
    lines = [l for l in re.split(r"\r\n|\n|\r", text) if l.strip()]
    if not lines:
        raise SampleError("The sample file is empty.")
    fmt, delim = _detect_format(lines)
    layout = Layout(
        source_name=filename, ext=ext, fmt=fmt, delimiter=delim or ",", encoding=enc,
        newline=newline, final_newline=text.endswith(("\n", "\r")),
    )
    if fmt == "delimited":
        parsed = [[v for v, _ in split_delimited(l, delim)] for l in lines]
        tokens = [p[0] for p in parsed]
        keys = [(len(p), None) for p in parsed]
    else:
        lens = Counter(len(l) for l in lines)
        len_ok = lens.most_common(1)[0][1] / len(lines) >= 0.6
        tokens = [l.split(" ", 1)[0] if l[1:2] == " " else None for l in lines]
        keys = [(len(l) if len_ok else None, None) for l in lines]
    # H/D/T style files: a record-type token that is identical on every middle line and differs at both ends
    if len(lines) >= 3 and None not in tokens[:1] + tokens[-1:] and len(set(tokens[1:-1])) == 1:
        if tokens[0] != tokens[1] and tokens[-1] != tokens[1] and None not in tokens:
            keys = [(k[0], t) for k, t in zip(keys, tokens)]
    h, t = _peel(keys)
    layout.header = lines[:h]
    layout.trailer = lines[len(lines) - t :]
    layout.details = lines[h : len(lines) - t]
    if not layout.details:
        raise SampleError("No detail records found in the sample file.")

    names: list[str] | None = None
    if fmt == "delimited":
        rows = [[v for v, _ in split_delimited(l, delim)] for l in layout.details]
        ncols = Counter(len(r) for r in rows).most_common(1)[0][0]
        hdr_fields = [v for v, _ in split_delimited(layout.header[-1], delim)] if layout.header else None
        if hdr_fields and len(hdr_fields) == ncols and _looks_like_header(hdr_fields, rows):
            layout.has_column_header, names = True, [c.strip() for c in hdr_fields]
        elif len(rows) > 1 and _looks_like_header(rows[0], rows[1:]):
            layout.header.append(layout.details.pop(0))
            layout.has_column_header, names = True, [c.strip() for c in rows[0]]
        quoted = [False] * ncols
        for l in layout.details:
            for i, (_, q) in enumerate(split_delimited(l, delim)[:ncols]):
                quoted[i] = quoted[i] or q
        layout.columns = [Column(name=f"Col{i + 1}", quoted=quoted[i]) for i in range(ncols)]
    else:
        def textual(l: str) -> bool:
            return not any(ch.isdigit() for ch in l)

        rest = lambda: any(ch.isdigit() for l in layout.details for ch in l)  # noqa: E731
        if layout.header and textual(layout.header[-1]) and rest():
            layout.has_column_header = True
        elif len(layout.details) > 1 and textual(layout.details[0]) and any(
            ch.isdigit() for l in layout.details[1:] for ch in l
        ):
            layout.header.append(layout.details.pop(0))
            layout.has_column_header = True
        layout.line_length = max(len(l) for l in layout.details)
        occupied = [False] * layout.line_length
        for l in layout.details:
            for i, ch in enumerate(l):
                if ch != " ":
                    occupied[i] = True
        cols, i = [], 0
        while i < len(occupied):
            if occupied[i]:
                j = i
                while j < len(occupied) and occupied[j]:
                    j += 1
                cols.append(Column(name=f"Col{len(cols) + 1}", start=i, width=j - i))
                i = j
            else:
                i += 1
        layout.columns = cols
        if layout.has_column_header:
            hdr = layout.header[-1].ljust(layout.line_length)
            names = [hdr[c.start : c.start + c.width].strip() for c in cols]
    if names:
        names = [n or f"Col{i + 1}" for i, n in enumerate(names[: len(layout.columns)])]
        names += [f"Col{i + 1}" for i in range(len(names), len(layout.columns))]
        for c, n in zip(layout.columns, _dedupe(names)):
            c.name = n
    return layout


def _parse_excel(filename: str, ext: str, data: bytes) -> Layout:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001 - surface any openpyxl failure to the user
        raise SampleError(f"Could not read the Excel sample: {exc}") from exc
    ws = wb.active
    cells = [list(r) for r in ws.iter_rows()]
    values = [[c.value for c in r] for r in cells]
    while values and all(_blank(v) for v in values[-1]):
        values.pop()
        cells.pop()
    if not values:
        raise SampleError("The sample workbook is empty.")
    ncols = max(len(r) for r in values)
    for r in values:
        r.extend([None] * (ncols - len(r)))

    def pattern(row: list[Any]) -> tuple[str, ...]:
        def kind(v: Any) -> str:
            if _blank(v):
                return "-"
            if isinstance(v, (datetime, date)):
                return "d"
            return "n" if _is_num(v) else "s"

        return tuple(kind(v) for v in row)

    def detail_like(pat: tuple[str, ...], modal: tuple[str, ...]) -> bool:
        filled = sum(k != "-" for k in pat)
        need = sum(k != "-" for k in modal)
        return filled >= 0.6 * need and all(a == b or a == "-" for a, b in zip(pat, modal))

    pats = [pattern(r) for r in values]
    middle = pats[1:-1] if len(pats) >= 3 else pats
    modal = Counter(middle).most_common(1)[0][0]
    flags = [detail_like(p, modal) for p in pats]
    first = flags.index(True) if True in flags else 0
    last = len(flags) - 1 - flags[::-1].index(True) if True in flags else len(flags) - 1
    h, t = (first, len(flags) - 1 - last) if len(pats) >= 3 else (0, 0)
    layout = Layout(source_name=filename, ext=ext, fmt="excel", sheet_name=ws.title)
    layout.header = values[:h]
    layout.trailer = values[len(values) - t :]
    layout.details = values[h : len(values) - t]
    if not layout.details:
        raise SampleError("No detail records found in the sample workbook.")
    names = None
    if layout.header and _looks_like_header(layout.header[-1], layout.details):
        names = [str(v).strip() if not _blank(v) else "" for v in layout.header[-1]]
        layout.has_column_header = True
    elif len(layout.details) > 1 and _looks_like_header(layout.details[0], layout.details[1:]):
        names = [str(v).strip() if not _blank(v) else "" for v in layout.details[0]]
        layout.header.append(layout.details.pop(0))
        layout.has_column_header = True
    layout.columns = [Column(name=names[i] if names and names[i] else f"Col{i + 1}") for i in range(ncols)]
    if names:
        for c, n in zip(layout.columns, _dedupe([c.name for c in layout.columns])):
            c.name = n

    def grab(c) -> dict:
        return {
            "font": copy.copy(c.font), "fill": copy.copy(c.fill), "border": copy.copy(c.border),
            "alignment": copy.copy(c.alignment), "number_format": c.number_format,
        }

    for ri in range(len(layout.header)):
        for ci, c in enumerate(cells[ri][:ncols]):
            layout.cell_styles[("header", ri, ci)] = grab(c)
    for ri in range(len(layout.trailer)):
        for ci, c in enumerate(cells[len(values) - t + ri][:ncols]):
            layout.cell_styles[("trailer", ri, ci)] = grab(c)
    for ci, c in enumerate(cells[len(layout.header)][:ncols]):
        layout.cell_styles[("detail", 0, ci)] = grab(c)
        layout.columns[ci].number_format = c.number_format
    for letter, dim in ws.column_dimensions.items():
        if dim.width:
            layout.col_widths[letter] = dim.width
    return layout


# --------------------------------------------------------------------------- analysis
def profile_column(col: Column, raw: list[Any]) -> None:
    """Fill in kind/format information of ``col`` from its sample values."""
    vals = [v for v in raw if not _blank(v)]
    col.blank_when_unused = len(vals) < len(raw)
    col.pool = list(dict.fromkeys(str(v).strip() if not isinstance(v, (datetime, date)) else str(v) for v in vals))[:40]
    col.native = False
    col.kind, col.decimals, col.zero_pad, col.thousands, col.date_fmt = "text", 0, 0, False, ""
    col.align = "left"
    if not vals:
        return
    if all(isinstance(v, (datetime, date)) for v in vals):
        col.kind, col.native = "date", True
        return
    if all(_is_num(v) for v in vals):
        col.native = True
        dec = max(max(-Decimal(str(v)).normalize().as_tuple().exponent, 0) for v in vals)
        m = re.search(r"\.(0+)", col.number_format or "")
        col.decimals = len(m.group(1)) if m else dec
        col.kind = "decimal" if col.decimals else "int"
        col.align = "right"
        return
    strs = [str(v).strip() for v in vals]
    fmt = detect_date_format(strs)
    if fmt:
        col.kind, col.date_fmt = "date", fmt
        col.date_upper = all(s == s.upper() for s in strs) and any(c.isalpha() for s in strs for c in s)
        return
    plain = [s.replace(",", "") for s in strs]
    if all(re.fullmatch(r"-?\d+", p) for p in plain):
        col.kind, col.align = "int", "right"
        if all(re.fullmatch(r"-?\d{1,3}(,\d{3})+", s) for s in strs if len(s) > 4) and any("," in s for s in strs):
            col.thousands = True
        lens = {len(s) for s in strs}
        if len(lens) == 1 and any(s.startswith("0") and len(s) > 1 for s in strs):
            col.zero_pad = lens.pop()
    elif all(re.fullmatch(r"-?\d+\.\d+", p) for p in plain):
        col.kind, col.align = "decimal", "right"
        col.decimals = max(len(p.split(".")[1]) for p in plain)
        col.thousands = any("," in s for s in strs)
        lens = {len(s) for s in strs}
        if len(lens) == 1 and any(s.startswith("0") for s in strs):
            col.zero_pad = lens.pop()
    # reference pattern: shared non-digit prefix + fixed number of digits, all unique
    m = [re.fullmatch(r"(\D*)(\d+)", s) for s in strs]
    if col.kind != "decimal" and len(strs) >= 3 and all(m) and len(set(strs)) == len(strs):
        prefixes = {x.group(1) for x in m}
        digits = {len(x.group(2)) for x in m}
        if len(prefixes) == 1 and len(digits) == 1 and (prefixes != {""} or digits.pop() >= 5):
            col.prefix, col.digits = prefixes.pop(), len(m[0].group(2))


_NAME_RULES: list[tuple[str, str]] = [
    ("manufacturer", r"manuf|mfr|mfg|supplier|dealer|vendor|brand|payee.?code|^mc$"),
    ("file_date", r"file.?date|creat|process|run.?date|generat|batch.?date|extract"),
    ("payment_date", r"date|due|settle|period"),
    ("debit_amount", r"debit|\bdr\b|^dr"),
    ("credit_amount", r"credit|\bcr\b|^cr"),
    ("amount", r"amount|amt|value|payment|total|net|gross|sum|pay"),
    ("sequence", r"seq|line|row|serial|record.?no|^no$|^#$"),
    ("reference", r"ref|invoice|inv|\bid\b|number|trans|batch"),
]


def _detect_role(col: Column, raw: list[Any], taken: set[str]) -> str:
    vals = [str(v).strip().upper() for v in raw if not _blank(v)]
    if vals and col.kind == "text" and set(vals) <= (_DEBIT_TOKENS | _CREDIT_TOKENS):
        return "dc_indicator"
    name = col.name.lower().strip()
    for role, rx in _NAME_RULES:
        if not re.search(rx, name):
            continue
        if role == "manufacturer" and col.kind in ("text", "int"):
            return role
        if role in ("file_date", "payment_date") and col.kind == "date":
            return role
        if role in ("amount", "debit_amount", "credit_amount") and col.kind in ("int", "decimal"):
            return role
        if role == "sequence" and col.kind == "int":
            return role
        if role == "reference" and col.kind in ("text", "int"):
            return role
    if col.kind == "date":
        return "payment_date" if "payment_date" not in taken else "file_date"
    if col.kind == "int" and len(vals) >= 3:
        nums = [int(v) for v in vals if re.fullmatch(r"\d+", v)]
        if len(nums) == len(vals) and all(b - a == 1 for a, b in zip(nums, nums[1:])):
            return "sequence"
    if col.kind == "int" and col.zero_pad >= 8 and "amount" not in taken:
        return "amount"  # unnamed zero-padded number, e.g. 000000150025 (minor units)
    if col.prefix or col.digits:
        return "reference"
    if col.kind == "decimal":
        return "amount" if "amount" not in taken else "other"
    return "other"


def analyze(layout: Layout) -> None:
    """(Re)profile all columns, detect roles still set to ``auto`` and the header/trailer bindings."""
    taken: set[str] = {c.role for c in layout.columns if c.role != "auto"}
    for i, col in enumerate(layout.columns):
        raw = layout.column_values(i)
        decimals, implied = col.decimals, col.implied_point
        profile_column(col, raw)
        if implied:  # user-declared implied decimal point wins over what the sample suggests
            col.decimals = decimals
        if col.role == "auto":
            col.role = _detect_role(col, raw, taken)
            taken.add(col.role)
    _fallback_manufacturer(layout)
    for i, col in enumerate(layout.columns):
        raw = layout.column_values(i)
        if col.role == "dc_indicator":
            _set_tokens(col, raw)
        if col.role in ("amount", "credit_amount", "debit_amount") and col.kind == "int" and col.zero_pad >= 6 and not col.implied_point:
            col.implied_point, col.decimals = True, 2  # e.g. 0000012345 == 123.45 (editable by the user)
    layout.debit_style = _debit_style(layout)
    layout.bindings = _find_bindings(layout)


def _fallback_manufacturer(layout: Layout) -> None:
    if layout.role_column("manufacturer"):
        return
    for i, col in enumerate(layout.columns):
        if col.role == "other" and col.kind == "text" and col.pool and all(
            re.fullmatch(r"[A-Z0-9]{2,8}", p) for p in col.pool
        ):
            col.role = "manufacturer"
            return


def _set_tokens(col: Column, raw: list[Any]) -> None:
    seen = {str(v).strip() for v in raw if not _blank(v)}
    cr = next((s for s in seen if s.upper() in _CREDIT_TOKENS), None)
    dr = next((s for s in seen if s.upper() in _DEBIT_TOKENS), None)
    pairs = {"C": "D", "CR": "DR", "CREDIT": "DEBIT", "CRD": "DB"}
    if cr and not dr:
        dr = pairs[cr.upper()]
        dr = dr.lower() if cr.islower() else dr
    if dr and not cr:
        inv = {v: k for k, v in pairs.items()}
        cr = inv.get(dr.upper(), "CR")
        cr = cr.lower() if dr.islower() else cr
    col.credit_token, col.debit_token = cr or "CR", dr or "DR"


def _debit_style(layout: Layout) -> str:
    roles = {c.role for c in layout.columns}
    if "dc_indicator" in roles:
        return "indicator"
    if {"credit_amount", "debit_amount"} <= roles:
        return "columns"
    amt = layout.role_column("amount")
    if amt is not None:
        for v in layout.column_values(layout.columns.index(amt)):
            d = to_decimal(v, amt)
            if d is not None and d < 0:
                return "signed"
    return "unsigned"


def sample_items(layout: Layout) -> list[tuple[Decimal, bool]]:
    """(absolute amount, is_debit) for every detail record of the sample."""
    cols = layout.columns
    idx = {c.role: i for i, c in reversed(list(enumerate(cols)))}
    items = []
    for rec in layout.details:
        f = layout.fields(rec)
        get = lambda role: to_decimal(f[idx[role]], cols[idx[role]]) if role in idx and idx[role] < len(f) else None  # noqa: E731
        if layout.debit_style == "columns":
            cr, dr = get("credit_amount"), get("debit_amount")
            if dr and dr != 0:
                items.append((abs(dr), True))
            else:
                items.append((abs(cr or Decimal(0)), False))
            continue
        amt = get("amount")
        if amt is None:
            continue
        is_debit = amt < 0
        if layout.debit_style == "indicator" and "dc_indicator" in idx:
            tok = str(f[idx["dc_indicator"]]).strip().upper()
            is_debit = tok in _DEBIT_TOKENS
        items.append((abs(amt), is_debit))
    return items


def compute_stats(items: list[tuple[Decimal, bool]]) -> dict[str, Decimal]:
    credit = sum((a for a, d in items if not d), Decimal(0))
    debit = sum((a for a, d in items if d), Decimal(0))
    return {
        "count": Decimal(len(items)),
        "credit_sum": credit,
        "debit_sum": debit,
        "total_sum": credit + debit,
        "net": credit - debit,
    }


def _find_bindings(layout: Layout) -> list[Binding]:
    stats = compute_stats(sample_items(layout))
    amt = next((c for c in layout.columns if c.role in ("amount", "credit_amount", "debit_amount")), None)
    dec = amt.decimals if amt else 2
    out: list[Binding] = []

    def match(token: str, native: Any = None) -> tuple[str, int, bool] | None:
        s = token.strip().replace(",", "")
        if native is not None and _is_num(native):
            s = str(native)
        if not re.fullmatch(r"-?\d+(\.\d+)?", s):
            return None
        d = Decimal(s)
        has_dot = "." in s
        for stat in ("count", "credit_sum", "debit_sum", "total_sum", "net"):
            target = stats[stat]
            if target == 0 or (stat == "count" and target < 2):
                continue
            if stat == "count":
                if not has_dot and d == target:
                    return stat, 0, has_dot
            elif has_dot:
                if d == target:
                    return stat, len(s.split(".")[1]), True
            elif d == target * (Decimal(10) ** dec) or (native is not None and d == target):
                return stat, dec, False
        return None

    for section, records in (("header", layout.header), ("trailer", layout.trailer)):
        for li, rec in enumerate(records):
            if section == "header" and layout.has_column_header and li == len(records) - 1:
                continue
            if layout.fmt == "excel":
                for ci, v in enumerate(rec):
                    m = match(str(v), v) if _is_num(v) else None
                    if m:
                        out.append(Binding(section, li, m[0], index=ci, decimals=m[1], has_dot=m[2]))
            elif layout.fmt == "delimited":
                for ci, (v, _) in enumerate(split_delimited(rec, layout.delimiter)):
                    m = match(v)
                    if m:
                        out.append(Binding(section, li, m[0], index=ci, decimals=m[1], has_dot=m[2],
                                           pad=len(v) if v.startswith("0") and len(v) > 1 else 0))
            else:
                for tok in re.finditer(r"\S+", rec):
                    m = match(tok.group(0))
                    if m:
                        v = tok.group(0)
                        out.append(Binding(section, li, m[0], start=tok.start(), end=tok.end(), decimals=m[1],
                                           has_dot=m[2], pad=len(v) if v.startswith("0") and len(v) > 1 else 0))
    return out


# --------------------------------------------------------------------------- user overrides
def apply_overrides(layout: Layout, edits: list[dict]) -> Layout:
    """Return a copy of ``layout`` with user-edited column settings applied.

    ``edits`` rows hold: name, role, start, width, decimals, implied_point.
    """
    new = copy.deepcopy(layout)
    cols: list[Column] = []
    for i, e in enumerate(edits):
        base = copy.deepcopy(layout.columns[i]) if i < len(layout.columns) else Column(name="")
        base.name = str(e.get("name") or f"Col{i + 1}")
        base.role = e.get("role") or "other"
        if layout.fmt == "fixed":
            base.start, base.width = int(e.get("start") or 0), max(int(e.get("width") or 1), 1)
        base.decimals = int(e.get("decimals") or 0)
        base.implied_point = bool(e.get("implied_point"))
        cols.append(base)
    new.columns = cols
    if new.fmt == "fixed":
        new.line_length = max([new.line_length] + [c.start + c.width for c in cols])
    analyze(new)
    return new
