"""Write payments in the exact layout of a sample file."""
from __future__ import annotations

import io
import random
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from typing import Any

from .generator import Payment
from .layout import Binding, Column, Layout, compute_stats, join_delimited, refresh_dates, split_delimited

ZERO = Decimal("0")


# ------------------------------------------------------------------ value formatting
def format_number(value: Decimal, col: Column) -> str:
    dec = col.decimals
    v = value.quantize(Decimal(1).scaleb(-dec), rounding=ROUND_HALF_UP)
    if v == 0:
        v = abs(v)
    neg = v < 0
    if col.implied_point:
        body = str(int((abs(v) * (Decimal(10) ** dec)).to_integral_value()))
    else:
        body = f"{abs(v):,.{dec}f}" if col.thousands else f"{abs(v):.{dec}f}"
    if col.zero_pad:
        body = body.zfill(col.zero_pad - (1 if neg else 0))
    return ("-" if neg else "") + body


def format_date(d: date, col: Column) -> str:
    out = d.strftime(col.date_fmt or "%Y%m%d")
    return out.upper() if col.date_upper else out


class _Context:
    def __init__(self, layout: Layout, file_date: date, rng: random.Random):
        self.layout, self.file_date, self.rng = layout, file_date, rng
        self.used_refs: dict[int, set[str]] = {}

    def reference(self, ci: int, col: Column) -> str:
        used = self.used_refs.setdefault(ci, set())
        digits = col.digits or 8
        while True:
            val = f"{col.prefix}{self.rng.randint(10 ** (digits - 1), 10**digits - 1)}"
            if val not in used:
                used.add(val)
                return val

    def cell(self, ci: int, col: Column, p: Payment, seq: int) -> Any:
        """Python-native value for a detail cell (Decimal / date / str / None)."""
        role, style = col.role, self.layout.debit_style
        if role == "manufacturer":
            return p.manufacturer
        if role == "payment_date":
            return p.pay_date
        if role == "file_date":
            return self.file_date
        if role == "amount":
            return -p.amount if p.is_debit and style == "signed" else p.amount
        if role in ("credit_amount", "debit_amount"):
            mine = p.is_debit == (role == "debit_amount")
            if mine:
                return p.amount
            return None if col.blank_when_unused else ZERO
        if role == "dc_indicator":
            return col.debit_token if p.is_debit else col.credit_token
        if role == "sequence":
            return seq
        if role == "reference" and (col.prefix or col.digits):
            return self.reference(ci, col)
        if col.pool:
            return self.rng.choice(col.pool) if len(col.pool) > 1 else col.pool[0]
        return ""

    def as_text(self, col: Column, value: Any) -> str:
        if value is None or value == "":
            return ""
        if isinstance(value, date):
            return format_date(value, col)
        if isinstance(value, Decimal):
            return format_number(value, col)
        if isinstance(value, int):
            return str(value).zfill(col.zero_pad) if col.zero_pad else str(value)
        return str(value)


def _stats_for(payments: list[Payment]) -> dict[str, Decimal]:
    return compute_stats([(p.amount, p.is_debit) for p in payments])


def _binding_text(b: Binding, stats: dict[str, Decimal]) -> str:
    value = stats[b.stat]
    if b.stat == "count":
        s = str(int(value))
    elif b.has_dot:
        s = f"{value:.{b.decimals}f}"
    else:
        s = str(int((value * (Decimal(10) ** b.decimals)).to_integral_value(ROUND_HALF_UP)))
    return s.zfill(b.pad) if b.pad else s


# ------------------------------------------------------------------ text output
def _render_text(layout: Layout, payments: list[Payment], file_date: date, rng: random.Random) -> bytes:
    ctx = _Context(layout, file_date, rng)
    stats = _stats_for(payments)
    lines: list[str] = []

    def side(section: str, records: list[str]) -> list[str]:
        out = []
        for li, rec in enumerate(records):
            if section == "header" and layout.has_column_header and li == len(records) - 1:
                out.append(rec)
                continue
            bs = [b for b in layout.bindings if b.section == section and b.line == li]
            if layout.fmt == "fixed":
                for b in sorted(bs, key=lambda b: -b.start):
                    width = b.end - b.start
                    rec = rec[: b.start] + _binding_text(b, stats).rjust(width) + rec[b.end :]
                out.append(refresh_dates(rec, file_date))
            else:
                rec = refresh_dates(rec, file_date)
                if bs:
                    parts = split_delimited(rec, layout.delimiter)
                    vals = [v for v, _ in parts]
                    for b in bs:
                        if b.index < len(vals):
                            vals[b.index] = _binding_text(b, stats)
                    rec = join_delimited(vals, [q for _, q in parts], layout.delimiter)
                out.append(rec)
        return out

    lines += side("header", layout.header)
    for seq, p in enumerate(payments, start=1):
        cells = [ctx.as_text(c, ctx.cell(i, c, p, seq)) for i, c in enumerate(layout.columns)]
        if layout.fmt == "fixed":
            buf = [" "] * layout.line_length
            for c, s in zip(layout.columns, cells):
                if len(s) > c.width:
                    s = s[-c.width :] if c.align == "right" else s[: c.width]
                pos = c.start + (c.width - len(s) if c.align == "right" else 0)
                buf[pos : pos + len(s)] = s
            lines.append("".join(buf))
        else:
            lines.append(join_delimited(cells, [c.quoted for c in layout.columns], layout.delimiter))
    lines += side("trailer", layout.trailer)
    text = layout.newline.join(lines) + (layout.newline if layout.final_newline else "")
    return text.encode(layout.encoding)


# ------------------------------------------------------------------ excel output
def _render_excel(layout: Layout, payments: list[Payment], file_date: date, rng: random.Random) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    ctx = _Context(layout, file_date, rng)
    stats = _stats_for(payments)
    wb = Workbook()
    ws = wb.active
    ws.title = layout.sheet_name
    row = 1

    def style(cell, key) -> None:
        st = layout.cell_styles.get(key)
        if st:
            cell.font, cell.fill, cell.border, cell.alignment = st["font"], st["fill"], st["border"], st["alignment"]
            cell.number_format = st["number_format"]

    def side(section: str, records: list[list[Any]]) -> None:
        nonlocal row
        for li, rec in enumerate(records):
            is_colhdr = section == "header" and layout.has_column_header and li == len(records) - 1
            bs = {b.index: b for b in layout.bindings if b.section == section and b.line == li}
            for ci, v in enumerate(rec):
                if not is_colhdr:
                    if ci in bs:
                        d = Decimal(_binding_text(bs[ci], stats))
                        v = int(d) if bs[ci].decimals == 0 and not bs[ci].has_dot else float(d)
                    elif isinstance(v, (datetime, date)):
                        v = datetime(file_date.year, file_date.month, file_date.day)
                    elif isinstance(v, str):
                        v = refresh_dates(v, file_date)
                cell = ws.cell(row=row, column=ci + 1, value=v)
                style(cell, (section, li, ci))
            row += 1

    side("header", layout.header)
    for seq, p in enumerate(payments, start=1):
        for ci, c in enumerate(layout.columns):
            v = ctx.cell(ci, c, p, seq)
            if c.native and isinstance(v, date):
                v = datetime(v.year, v.month, v.day)
            elif c.native and isinstance(v, Decimal):
                v = int(v) if c.decimals == 0 else float(v.quantize(Decimal(1).scaleb(-c.decimals)))
            elif c.native and v is None:
                v = None
            else:
                v = ctx.as_text(c, v) or None
            cell = ws.cell(row=row, column=ci + 1, value=v)
            style(cell, ("detail", 0, ci))
            if c.number_format:
                cell.number_format = c.number_format
        row += 1
    side("trailer", layout.trailer)
    for letter, width in layout.col_widths.items():
        ws.column_dimensions[letter].width = width
    for ci in range(len(layout.columns)):
        letter = get_column_letter(ci + 1)
        if letter not in layout.col_widths:
            ws.column_dimensions[letter].width = max(12, len(layout.columns[ci].name) + 2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(layout: Layout, payments: list[Payment], file_date: date, seed: int | None = None) -> bytes:
    rng = random.Random(seed)
    if layout.fmt == "excel":
        return _render_excel(layout, payments, file_date, rng)
    return _render_text(layout, payments, file_date, rng)
